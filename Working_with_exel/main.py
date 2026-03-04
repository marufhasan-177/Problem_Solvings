import openpyxl as xl
wb = xl.load_workbook(r'D:\All_Codes\Python\Working_with_exel\transections.xlsx')
sheet = wb['Sheet1']
cell = sheet['a2']

for i in range(3,sheet.max_row+1):
    pnt = sheet.cell(i,2)
    print(pnt.value)